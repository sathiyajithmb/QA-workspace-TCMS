import os, hmac, hashlib, secrets, io, csv, json, re, logging, time, hashlib as _hashlib, xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from typing import Optional
from contextlib import asynccontextmanager
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query, Request, Response, Body
from fastapi.encoders import jsonable_encoder
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, Column, Integer, String, Text, ForeignKey, DateTime, Boolean, JSON, select, func, UniqueConstraint, Index, inspect, text
from sqlalchemy.orm import declarative_base, sessionmaker, Session
import jwt

DATABASE_URL=os.getenv('DATABASE_URL','sqlite:///./dev.db')
SECRET=os.getenv('JWT_SECRET','development-only-change-me')
engine=create_engine(DATABASE_URL, pool_pre_ping=True, connect_args={'check_same_thread':False} if DATABASE_URL.startswith('sqlite') else {})
SessionLocal=sessionmaker(bind=engine)
Base=declarative_base()
UPLOAD_DIR=os.getenv('UPLOAD_DIR',os.path.join(os.path.dirname(__file__),'uploads'))
MAX_UPLOAD_BYTES=int(os.getenv('MAX_UPLOAD_BYTES','10485760'))
SESSION_HOURS=int(os.getenv('SESSION_HOURS','8'))
COOKIE_SECURE=os.getenv('COOKIE_SECURE','false').lower()=='true'
def now(): return datetime.now(timezone.utc)
class User(Base):
    __tablename__='users'
    id=Column(Integer,primary_key=True); email=Column(String(255),unique=True,nullable=False); name=Column(String(120),nullable=False); password_hash=Column(String(300),nullable=False); role=Column(String(20),default='tester'); roles=Column(JSON,nullable=False,default=list); active=Column(Boolean,default=True)
class Project(Base):
    __tablename__='projects'
    id=Column(Integer,primary_key=True); key=Column(String(40),unique=True); name=Column(String(180),nullable=False); description=Column(Text,default=''); archived=Column(Boolean,default=False,nullable=False); created_at=Column(DateTime(timezone=True),default=now)
class Membership(Base):
    __tablename__='memberships'; __table_args__=(UniqueConstraint('project_id','user_id'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); role=Column(String(20),default='tester',nullable=False)
class TestFolder(Base):
    __tablename__='test_folders'; __table_args__=(UniqueConstraint('project_id','name'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); parent_id=Column(Integer,ForeignKey('test_folders.id'),nullable=True); name=Column(String(180),nullable=False)
class TestCase(Base):
    __tablename__='test_cases'
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); folder_id=Column(Integer,ForeignKey('test_folders.id'),nullable=True); title=Column(String(300),nullable=False); priority=Column(String(20),default='Medium'); kind=Column(String(20),default='Manual'); status=Column(String(30),default='Active'); version=Column(Integer,default=1); archived=Column(Boolean,default=False); author_id=Column(Integer,ForeignKey('users.id'),nullable=True); created_at=Column(DateTime(timezone=True),default=now); updated_at=Column(DateTime(timezone=True),default=now,onupdate=now); custom_values=Column(JSON,default=dict)
class CaseVersion(Base):
    __tablename__='case_versions'; __table_args__=(UniqueConstraint('case_id','version'),)
    id=Column(Integer,primary_key=True); case_id=Column(Integer,ForeignKey('test_cases.id'),nullable=False); version=Column(Integer,nullable=False); preconditions=Column(Text,default=''); steps=Column(JSON,default=list); author_id=Column(Integer,ForeignKey('users.id'),nullable=True); created_at=Column(DateTime(timezone=True),default=now)
class TestPlan(Base):
    __tablename__='test_plans'
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); name=Column(String(200),nullable=False); release=Column(String(100),default=''); scope=Column(Text,default=''); owner_id=Column(Integer,ForeignKey('users.id'),nullable=True); archived=Column(Boolean,default=False)
class PlanCase(Base):
    __tablename__='plan_cases'; __table_args__=(UniqueConstraint('plan_id','case_id'),)
    id=Column(Integer,primary_key=True); plan_id=Column(Integer,ForeignKey('test_plans.id'),nullable=False); case_id=Column(Integer,ForeignKey('test_cases.id'),nullable=False)
class Cycle(Base):
    __tablename__='cycles'
    id=Column(Integer,primary_key=True); plan_id=Column(Integer,ForeignKey('test_plans.id'),nullable=False); name=Column(String(200),nullable=False); environment=Column(String(100),default='QA'); start_date=Column(DateTime(timezone=True),nullable=True); end_date=Column(DateTime(timezone=True),nullable=True)
class CycleCase(Base):
    __tablename__='cycle_cases'; __table_args__=(UniqueConstraint('cycle_id','case_id'),)
    id=Column(Integer,primary_key=True); cycle_id=Column(Integer,ForeignKey('cycles.id'),nullable=False); case_id=Column(Integer,ForeignKey('test_cases.id'),nullable=False); version_id=Column(Integer,ForeignKey('case_versions.id'),nullable=False); assigned_to=Column(Integer,ForeignKey('users.id'),nullable=True)
class Execution(Base):
    __tablename__='executions'
    id=Column(Integer,primary_key=True); cycle_case_id=Column(Integer,ForeignKey('cycle_cases.id'),nullable=False); user_id=Column(Integer,ForeignKey('users.id'),nullable=True); status=Column(String(20),nullable=False); source=Column(String(20),default='Manual'); notes=Column(Text,default=''); actual_result=Column(Text,default=''); step_results=Column(JSON,default=list); duration=Column(String(30),default=''); created_at=Column(DateTime(timezone=True),default=now)
class Defect(Base):
    __tablename__='defects'
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); title=Column(String(300),nullable=False); description=Column(Text,default=''); severity=Column(String(20),default='Medium'); priority=Column(String(20),default='Medium'); status=Column(String(30),default='New'); assigned_to=Column(Integer,ForeignKey('users.id'),nullable=True); execution_id=Column(Integer,ForeignKey('executions.id'),nullable=True); created_at=Column(DateTime(timezone=True),default=now); updated_at=Column(DateTime(timezone=True),default=now,onupdate=now)
class DefectExecution(Base):
    __tablename__='defect_executions'; __table_args__=(UniqueConstraint('defect_id','execution_id'),)
    id=Column(Integer,primary_key=True); defect_id=Column(Integer,ForeignKey('defects.id'),nullable=False); execution_id=Column(Integer,ForeignKey('executions.id'),nullable=False)
class AutomationMap(Base):
    __tablename__='automation_maps'; __table_args__=(UniqueConstraint('project_id','external_id'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); external_id=Column(String(400),nullable=False); case_id=Column(Integer,ForeignKey('test_cases.id'),nullable=False)
class ImportLog(Base):
    __tablename__='import_logs'; __table_args__=(UniqueConstraint('cycle_id','checksum'),)
    id=Column(Integer,primary_key=True); cycle_id=Column(Integer,ForeignKey('cycles.id'),nullable=False); filename=Column(String(255)); checksum=Column(String(64),nullable=True); imported=Column(Integer,default=0); unmatched=Column(JSON,default=list); errors=Column(JSON,default=list); created_at=Column(DateTime(timezone=True),default=now)
class Attachment(Base):
    __tablename__='attachments'
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); execution_id=Column(Integer,ForeignKey('executions.id'),nullable=True); filename=Column(String(255),nullable=False); content_type=Column(String(100),nullable=False); size=Column(Integer,nullable=False); storage_key=Column(String(80),unique=True,nullable=False); uploaded_by=Column(Integer,ForeignKey('users.id'),nullable=False); created_at=Column(DateTime(timezone=True),default=now)
class AuditEvent(Base):
    __tablename__='audit_events'; __table_args__=(Index('ix_audit_project_created','project_id','created_at'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=True); actor_id=Column(Integer,ForeignKey('users.id'),nullable=True); action=Column(String(80),nullable=False); entity_type=Column(String(80),nullable=False); entity_id=Column(String(80),nullable=False); before=Column(JSON,nullable=True); after=Column(JSON,nullable=True); created_at=Column(DateTime(timezone=True),default=now)
class ApiToken(Base):
    __tablename__='api_tokens'
    id=Column(Integer,primary_key=True); name=Column(String(120),nullable=False); token_hash=Column(String(64),unique=True,nullable=False); prefix=Column(String(12),nullable=False); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); expires_at=Column(DateTime(timezone=True),nullable=True); revoked=Column(Boolean,default=False); created_at=Column(DateTime(timezone=True),default=now)
class Notification(Base):
    __tablename__='notifications'; __table_args__=(Index('ix_notifications_user_created','user_id','created_at'),)
    id=Column(Integer,primary_key=True); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); title=Column(String(200),nullable=False); message=Column(Text,nullable=False); link=Column(String(300),default=''); read=Column(Boolean,default=False); created_at=Column(DateTime(timezone=True),default=now)
class RequirementRecord(Base):
    __tablename__='requirements'; __table_args__=(UniqueConstraint('project_id','key'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); key=Column(String(80),nullable=False); title=Column(String(300),nullable=False); description=Column(Text,default=''); status=Column(String(30),default='Active'); created_at=Column(DateTime(timezone=True),default=now)
class RequirementCase(Base):
    __tablename__='requirement_cases'; __table_args__=(UniqueConstraint('requirement_id','case_id'),)
    id=Column(Integer,primary_key=True); requirement_id=Column(Integer,ForeignKey('requirements.id'),nullable=False); case_id=Column(Integer,ForeignKey('test_cases.id'),nullable=False)
class CustomField(Base):
    __tablename__='custom_fields'; __table_args__=(UniqueConstraint('project_id','name'),)
    id=Column(Integer,primary_key=True); project_id=Column(Integer,ForeignKey('projects.id'),nullable=False); name=Column(String(100),nullable=False); field_type=Column(String(20),nullable=False); required=Column(Boolean,default=False); options=Column(JSON,default=list)

def password_hash(password):
    salt=secrets.token_hex(16); digest=hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt),260000).hex(); return salt+':'+digest
def verify_password(password, stored):
    try:
        salt,digest=stored.split(':'); actual=hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt),260000).hex(); return hmac.compare_digest(actual,digest)
    except (ValueError,TypeError): return False

def db():
    with SessionLocal() as session: yield session
bearer=HTTPBearer(auto_error=False)
def user(request:Request,credentials:Optional[HTTPAuthorizationCredentials]=Depends(bearer),session:Session=Depends(db)):
    token=credentials.credentials if credentials else request.cookies.get('qa_session')
    if not token: raise HTTPException(401,'Authentication required')
    try: payload=jwt.decode(token,SECRET,algorithms=['HS256']); uid=int(payload['sub'])
    except Exception: raise HTTPException(401,'Invalid or expired token')
    u=session.get(User,uid)
    if not u or not u.active: raise HTTPException(401,'Inactive user')
    return u
def elevated(u):
    if not has_any_role(u,'admin','qa_lead','test_lead'): raise HTTPException(403,'Lead or admin required')
def admin(u):
    if not has_any_role(u,'admin'): raise HTTPException(403,'Admin required')
def accessible(session,u,pid):
    if has_any_role(u,'admin','qa_lead','test_lead'): return
    if not session.scalar(select(Membership).where(Membership.project_id==pid,Membership.user_id==u.id)): raise HTTPException(403,'Project access denied')
def obj(session,cls,id):
    found=session.get(cls,id)
    if not found: raise HTTPException(404,f'{cls.__name__} not found')
    return found
def project_for_cycle(session,cycle_id):
    cycle=obj(session,Cycle,cycle_id); plan=obj(session,TestPlan,cycle.plan_id); return plan.project_id

def data(o): return {c.name:getattr(o,c.name) for c in o.__table__.columns}

def ensure_database_schema(bind=None):
    target=bind or engine
    Base.metadata.create_all(target)
    additions={
        'users': {'roles':"JSON NOT NULL DEFAULT '[]'"},
        'projects': {'key':"VARCHAR(40)",'archived':"BOOLEAN NOT NULL DEFAULT FALSE"},
        'memberships': {'role':"VARCHAR(20) NOT NULL DEFAULT 'tester'"},
        'test_cases': {'folder_id':"INTEGER",'status':"VARCHAR(30) NOT NULL DEFAULT 'Active'",'author_id':"INTEGER",'created_at':"DATETIME",'updated_at':"DATETIME",'custom_values':"JSON NOT NULL DEFAULT '{}'"},
        'case_versions': {'author_id':"INTEGER"},
        'test_plans': {'release':"VARCHAR(100) NOT NULL DEFAULT ''",'scope':"TEXT NOT NULL DEFAULT ''",'owner_id':"INTEGER",'archived':"BOOLEAN NOT NULL DEFAULT FALSE"},
        'cycles': {'start_date':"DATETIME",'end_date':"DATETIME"},
        'executions': {'actual_result':"TEXT NOT NULL DEFAULT ''",'step_results':"JSON NOT NULL DEFAULT '[]'"},
        'defects': {'priority':"VARCHAR(20) NOT NULL DEFAULT 'Medium'",'created_at':"DATETIME",'updated_at':"DATETIME"},
        'import_logs': {'checksum':"VARCHAR(64)",'errors':"JSON NOT NULL DEFAULT '[]'"},
    }
    def migrate_columns(connection):
        for table, columns in additions.items():
            if table not in inspect(connection).get_table_names():
                continue
            present={column['name'] for column in inspect(connection).get_columns(table)}
            for name, declaration in columns.items():
                if name not in present:
                    connection.execute(text(f'ALTER TABLE {table} ADD COLUMN {name} {declaration}'))
        if 'projects' in inspect(connection).get_table_names():
            connection.execute(text("UPDATE projects SET key = 'PRJ-' || id WHERE key IS NULL OR key = ''"))
            connection.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS uq_projects_key ON projects (key)"))
        if 'memberships' in inspect(connection).get_table_names():
            connection.execute(text("UPDATE memberships SET role = COALESCE((SELECT role FROM users WHERE users.id = memberships.user_id), 'tester') WHERE role IS NULL OR role = ''"))
        if 'users' in inspect(connection).get_table_names():
            for user_id, legacy_role, stored_roles in connection.execute(select(User.id,User.role,User.roles)).all():
                if not stored_roles:
                    migrated_roles=['qa_lead' if legacy_role=='lead' else legacy_role or 'tester']
                    connection.execute(User.__table__.update().where(User.id==user_id).values(roles=migrated_roles))
    if bind is None:
        with engine.begin() as connection: migrate_columns(connection)
    else:
        migrate_columns(bind)
    Base.metadata.create_all(target)

def audit(session, actor_id, action, entity_type, entity_id, project_id=None, before=None, after=None):
    session.add(AuditEvent(project_id=project_id,actor_id=actor_id,action=action,entity_type=entity_type,entity_id=str(entity_id),before=jsonable_encoder(before) if before is not None else None,after=jsonable_encoder(after) if after is not None else None))

def notify(session, user_id, title, message, link=''):
    if user_id:
        session.add(Notification(user_id=user_id,title=title,message=message,link=link))

def project_role(session, u, project_id):
    if has_any_role(u,'admin'): return 'admin'
    if has_any_role(u,'qa_lead','test_lead'): return 'lead'
    membership=session.scalar(select(Membership).where(Membership.project_id==project_id,Membership.user_id==u.id))
    return membership.role if membership else None

GLOBAL_ROLES={'admin','qa_lead','test_lead','tester','viewer'}
ROLE_ALIASES={'lead':'qa_lead'}
def normalize_roles(roles):
    if not isinstance(roles,list) or not roles:
        raise HTTPException(400,'Assign at least one global role')
    normalized=[]
    for role in roles:
        value=ROLE_ALIASES.get(str(role).strip().lower(),str(role).strip().lower())
        if value not in GLOBAL_ROLES:
            raise HTTPException(400,f'Invalid global role: {role}')
        if value not in normalized:
            normalized.append(value)
    return normalized
def user_roles(u):
    roles=getattr(u,'roles',None)
    if roles:
        return [ROLE_ALIASES.get(role,role) for role in roles]
    legacy_role=getattr(u,'role','tester') or 'tester'
    return [ROLE_ALIASES.get(legacy_role,legacy_role)]
def has_any_role(u,*roles):
    return bool(set(user_roles(u)).intersection(ROLE_ALIASES.get(role,role) for role in roles))
def legacy_primary_role(roles):
    if 'admin' in roles: return 'admin'
    if 'qa_lead' in roles or 'test_lead' in roles: return 'lead'
    if 'tester' in roles: return 'tester'
    return 'viewer'

def require_project_role(session, u, project_id, allowed):
    if project_role(session,u,project_id) not in allowed:
        raise HTTPException(403,'Insufficient project role')

def validate_custom_values(session, project_id, values):
    definitions=session.scalars(select(CustomField).where(CustomField.project_id==project_id)).all()
    known={field.name for field in definitions}
    if set(values)-known: raise HTTPException(400,'Unknown custom field')
    for definition in definitions:
        value=values.get(definition.name)
        if definition.required and (value is None or value==''):
            raise HTTPException(400,f'{definition.name} is required')
        if value is None or value=='': continue
        if definition.field_type=='select' and value not in (definition.options or []):
            raise HTTPException(400,f'Invalid value for {definition.name}')
        if definition.field_type=='number' and (isinstance(value,bool) or not isinstance(value,(int,float))):
            raise HTTPException(400,f'{definition.name} must be a number')
        if definition.field_type=='boolean' and not isinstance(value,bool):
            raise HTTPException(400,f'{definition.name} must be true or false')

@asynccontextmanager
async def lifespan(app):
    from alembic import command
    from alembic.config import Config
    alembic_config=Config(os.path.join(os.path.dirname(__file__),'..','alembic.ini'))
    alembic_config.set_main_option('sqlalchemy.url',DATABASE_URL.replace('%','%%'))
    command.upgrade(alembic_config,'head')
    with SessionLocal() as s:
        if not s.scalar(select(func.count()).select_from(User)):
            email=os.getenv('ADMIN_EMAIL','admin@example.local'); pw=os.getenv('ADMIN_PASSWORD')
            if not pw or len(pw)<12: raise RuntimeError('Set ADMIN_PASSWORD with at least 12 characters')
            s.add(User(email=email,name='Administrator',password_hash=password_hash(pw),role='admin',roles=['admin'])); s.commit()
    os.makedirs(UPLOAD_DIR,exist_ok=True)
    yield
app=FastAPI(title='Docpanel QA API',lifespan=lifespan)
app.add_middleware(CORSMiddleware,allow_origins=[x.strip() for x in os.getenv('CORS_ORIGINS','http://localhost:5173,http://localhost:8080').split(',')],allow_credentials=True,allow_methods=['*'],allow_headers=['*'])
logger=logging.getLogger('qa_manager')
@app.middleware('http')
async def request_logging(request:Request,call_next):
    started=time.perf_counter()
    response=await call_next(request)
    response.headers['X-API-Version']='v1'
    logger.info(json.dumps({'event':'http_request','method':request.method,'path':request.url.path,'status':response.status_code,'duration_ms':round((time.perf_counter()-started)*1000,2)}))
    return response
@app.middleware('http')
async def versioned_api(request:Request,call_next):
    path=request.scope['path']
    if path=='/api/v1': request.scope['path']='/api'
    elif path.startswith('/api/v1/'): request.scope['path']='/api'+path[len('/api/v1'):]
    return await call_next(request)
class Login(BaseModel): email:str; password:str
class UserIn(BaseModel): email:str; name:str; password:str=Field(min_length=12); role:Optional[str]=None; roles:Optional[list[str]]=None
class UserUpdate(BaseModel): roles:Optional[list[str]]=None; active:Optional[bool]=None
class ProjectIn(BaseModel): key:str; name:str; description:str=''
class CaseIn(BaseModel): project_id:int; title:str; priority:str='Medium'; kind:str='Manual'; status:str='Active'; preconditions:str=''; steps:list[dict]=Field(default_factory=list); folder_id:Optional[int]=None; custom_values:dict=Field(default_factory=dict)
class PlanIn(BaseModel): project_id:int; name:str; release:str=''; scope:str=''; owner_id:Optional[int]=None; case_ids:list[int]=Field(default_factory=list)
class CycleIn(BaseModel): plan_id:int; name:str; environment:str='QA'; start_date:Optional[datetime]=None; end_date:Optional[datetime]=None
class AssignIn(BaseModel): case_id:int; assigned_to:Optional[int]=None; version_id:Optional[int]=None
class ExecuteIn(BaseModel): status:str; notes:str=''; actual_result:str=''; step_results:list[dict]=Field(default_factory=list)
class DefectIn(BaseModel): project_id:int; title:str; description:str=''; severity:str='Medium'; priority:str='Medium'; assigned_to:Optional[int]=None; execution_ids:list[int]=Field(default_factory=list); execution_id:Optional[int]=None
class MapIn(BaseModel): project_id:int; external_id:str; case_id:int
class MembershipIn(BaseModel): user_id:int; role:str='tester'
class RequirementIn(BaseModel): project_id:int; key:str; title:str; description:str=''; status:str='Active'
@app.post('/api/auth/login')
def login(body:Login,response:Response,s:Session=Depends(db)):
    u=s.scalar(select(User).where(User.email==body.email.lower()))
    if not u or not u.active or not verify_password(body.password, u.password_hash):
        raise HTTPException(
            status_code=401,
            detail="Invalid credentials"
        )
    token=jwt.encode({'sub':str(u.id),'exp':now()+timedelta(hours=SESSION_HOURS)},SECRET,algorithm='HS256')
    response.set_cookie('qa_session',token,max_age=SESSION_HOURS*3600,httponly=True,secure=COOKIE_SECURE,samesite='strict',path='/')
    return {'user':{'id':u.id,'email':u.email,'name':u.name,'role':u.role,'roles':user_roles(u)}}
@app.post('/api/auth/logout')
def logout(response:Response): response.delete_cookie('qa_session',path='/'); return {'ok':True}
@app.get('/api/me')
def me(u:User=Depends(user),s:Session=Depends(db)):
    return {'id':u.id,'email':u.email,'name':u.name,'role':u.role,'roles':user_roles(u),'project_roles':{m.project_id:m.role for m in s.scalars(select(Membership).where(Membership.user_id==u.id)).all()}}
@app.get('/api/users')
def users(project_id:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
    if has_any_role(u,'admin'): query=select(User)
    elif project_id is not None:
        accessible(s,u,project_id);query=select(User).join(Membership).where(Membership.project_id==project_id)
    else: raise HTTPException(403,'Project ID is required')
    return [{'id':x.id,'name':x.name,'email':x.email,'role':legacy_primary_role(user_roles(x)),'roles':user_roles(x),'active':x.active} for x in s.scalars(query).all()]
@app.post('/api/users')
def create_user(b:UserIn,s:Session=Depends(db),u:User=Depends(user)):
    admin(u)
    if s.scalar(select(User).where(User.email==b.email.lower())): raise HTTPException(409,'Email already exists')
    roles=normalize_roles(b.roles if b.roles is not None else [b.role or 'tester'])
    x=User(email=b.email.lower(),name=b.name,password_hash=password_hash(b.password),role=legacy_primary_role(roles),roles=roles); s.add(x);s.flush();audit(s,u.id,'create','user',x.id,after={'email':x.email,'roles':roles});s.commit();return {'id':x.id}
@app.patch('/api/users/{id}')
def update_user(id:int,b:Optional[UserUpdate]=Body(default=None),role:Optional[str]=None,active:Optional[bool]=None,s:Session=Depends(db),u:User=Depends(user)):
    admin(u); x=obj(s,User,id)
    before={'roles':user_roles(x),'active':x.active}
    next_roles=b.roles if b and b.roles is not None else ([role] if role is not None else None)
    next_active=b.active if b and b.active is not None else active
    if next_roles is not None:
        normalized_roles=normalize_roles(next_roles)
        if x.id==u.id and has_any_role(x,'admin') and 'admin' not in normalized_roles:
            other_admins=[candidate for candidate in s.scalars(select(User).where(User.active==True)).all() if candidate.id!=u.id and has_any_role(candidate,'admin')]
            if not other_admins: raise HTTPException(400,'Cannot remove the last active administrator')
        x.roles=normalized_roles
        x.role=legacy_primary_role(x.roles)
    if next_active is not None:
        if x.id==u.id and not next_active: raise HTTPException(400,'Cannot deactivate self')
        x.active=next_active
    audit(s,u.id,'update','user',x.id,before=before,after={'roles':user_roles(x),'active':x.active});s.commit();return {'ok':True,'roles':user_roles(x)}
@app.get('/api/projects')
def projects(include_archived:bool=False,s:Session=Depends(db),u:User=Depends(user)):
    q=select(Project) if has_any_role(u,'admin','qa_lead','test_lead') else select(Project).join(Membership).where(Membership.user_id==u.id)
    if not include_archived: q=q.where(Project.archived==False)
    return [data(x)|{'role':project_role(s,u,x.id)} for x in s.scalars(q).all()]
@app.post('/api/projects')
def create_project(b:ProjectIn,s:Session=Depends(db),u:User=Depends(user)):
    elevated(u)
    key=b.key.strip().upper()
    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_-]{1,39}',key): raise HTTPException(400,'Project key must be 2-40 letters, numbers, _ or -')
    if s.scalar(select(Project).where(Project.key==key)): raise HTTPException(409,'Project key already exists')
    x=Project(key=key,name=b.name,description=b.description);s.add(x);s.flush();s.add(Membership(project_id=x.id,user_id=u.id,role='admin' if has_any_role(u,'admin') else 'lead'));audit(s,u.id,'create','project',x.id,x.id,after={'key':x.key,'name':x.name});s.commit();return data(x)
@app.patch('/api/projects/{id}')
def update_project(id:int,b:ProjectIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id);require_project_role(s,u,id,('admin','lead'));x=obj(s,Project,id);before=data(x);key=b.key.strip().upper()
    if s.scalar(select(Project).where(Project.key==key,Project.id!=id)): raise HTTPException(409,'Project key already exists')
    x.key=key;x.name=b.name;x.description=b.description;audit(s,u.id,'update','project',id,id,before,b.model_dump());s.commit();return data(x)
@app.post('/api/projects/{id}/archive')
def archive_project(id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id);require_project_role(s,u,id,('admin','lead'));x=obj(s,Project,id);x.archived=True;audit(s,u.id,'archive','project',id,id,after=data(x));s.commit();return {'ok':True}
@app.post('/api/projects/{id}/restore')
def restore_project(id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id);require_project_role(s,u,id,('admin','lead'));x=obj(s,Project,id);x.archived=False;audit(s,u.id,'restore','project',id,id,after=data(x));s.commit();return {'ok':True}
@app.post('/api/projects/{id}/members')
def add_member(id:int,b:MembershipIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id);require_project_role(s,u,id,('admin','lead'));obj(s,Project,id);obj(s,User,b.user_id)
    if b.role not in ('admin','lead','tester','viewer'): raise HTTPException(400,'Invalid project role')
    membership=s.scalar(select(Membership).where(Membership.project_id==id,Membership.user_id==b.user_id))
    if membership: membership.role=b.role
    else: s.add(Membership(project_id=id,user_id=b.user_id,role=b.role))
    notify(s,b.user_id,'Project access granted',f'You were added to project #{id} as {b.role}.',f'/projects/{id}')
    audit(s,u.id,'grant_access','membership',b.user_id,id,after={'role':b.role});s.commit()
    return {'ok':True}
@app.get('/api/projects/{id}/members')
def project_members(id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id)
    return [{'user_id':m.user_id,'name':obj(s,User,m.user_id).name,'email':obj(s,User,m.user_id).email,'role':m.role} for m in s.scalars(select(Membership).where(Membership.project_id==id)).all()]
@app.get('/api/projects/{id}/available-users')
def available_project_users(id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,id);require_project_role(s,u,id,('admin','lead'))
    assigned=select(Membership.user_id).where(Membership.project_id==id)
    return [{'id':x.id,'name':x.name,'email':x.email,'role':legacy_primary_role(user_roles(x)),'roles':user_roles(x),'active':x.active} for x in s.scalars(select(User).where(User.active==True,User.id.not_in(assigned))).all()]
@app.get('/api/cases')
def cases(project_id:int,q:str='',priority:Optional[str]=None,kind:Optional[str]=None,folder_id:Optional[int]=None,include_archived:bool=False,offset:int=Query(0,ge=0),limit:int=Query(100,ge=1,le=500),s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);query=select(TestCase).where(TestCase.project_id==project_id)
    if not include_archived: query=query.where(TestCase.archived==False)
    if q: query=query.where(TestCase.title.ilike(f'%{q}%'))
    if priority: query=query.where(TestCase.priority==priority)
    if kind: query=query.where(TestCase.kind==kind)
    if folder_id is not None: query=query.where(TestCase.folder_id==folder_id)
    return [data(x) for x in s.scalars(query.order_by(TestCase.id.desc()).offset(offset).limit(limit)).all()]
@app.get('/api/cases/{id}')
def case(id:int,version:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,TestCase,id);accessible(s,u,x.project_id);v=s.scalar(select(CaseVersion).where(CaseVersion.case_id==id,CaseVersion.version==version if version else x.version))
    if not v: raise HTTPException(404,'Test case version not found')
    return data(x)|{'preconditions':v.preconditions,'steps':v.steps,'version_number':v.version}
@app.get('/api/cases/{id}/versions')
def case_versions(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,TestCase,id);accessible(s,u,x.project_id)
    return [data(v) for v in s.scalars(select(CaseVersion).where(CaseVersion.case_id==id).order_by(CaseVersion.version.desc())).all()]
@app.post('/api/cases')
def create_case(b:CaseIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead','tester'))
    validate_custom_values(s,b.project_id,b.custom_values)
    if b.priority not in ('Low','Medium','High','Critical') or b.kind not in ('Manual','Automated','Both') or b.status not in ('Active','Draft','Deprecated'): raise HTTPException(400,'Invalid test case metadata')
    if b.folder_id: 
        folder=obj(s,TestFolder,b.folder_id)
        if folder.project_id!=b.project_id: raise HTTPException(400,'Folder belongs to another project')
    x=TestCase(project_id=b.project_id,folder_id=b.folder_id,title=b.title,priority=b.priority,kind=b.kind,status=b.status,author_id=u.id,custom_values=b.custom_values);s.add(x);s.flush();s.add(CaseVersion(case_id=x.id,version=1,preconditions=b.preconditions,steps=b.steps,author_id=u.id));audit(s,u.id,'create','test_case',x.id,b.project_id,after=data(x));s.commit();return data(x)
@app.put('/api/cases/{id}')
def edit_case(id:int,b:CaseIn,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));before=data(x)
    if b.project_id!=x.project_id: raise HTTPException(400,'Cannot move case across projects')
    validate_custom_values(s,x.project_id,b.custom_values)
    if b.priority not in ('Low','Medium','High','Critical') or b.kind not in ('Manual','Automated','Both') or b.status not in ('Active','Draft','Deprecated'): raise HTTPException(400,'Invalid test case metadata')
    if b.folder_id:
        folder=obj(s,TestFolder,b.folder_id)
        if folder.project_id!=x.project_id: raise HTTPException(400,'Folder belongs to another project')
    x.title=b.title;x.priority=b.priority;x.kind=b.kind;x.status=b.status;x.folder_id=b.folder_id;x.custom_values=b.custom_values;x.version+=1;x.updated_at=now();s.add(CaseVersion(case_id=id,version=x.version,preconditions=b.preconditions,steps=b.steps,author_id=u.id));audit(s,u.id,'update','test_case',id,x.project_id,before,data(x));s.commit();return data(x)
@app.post('/api/cases/{id}/archive')
def archive_case(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));x.archived=True;audit(s,u.id,'archive','test_case',id,x.project_id,after=data(x));s.commit();return {'ok':True}
@app.post('/api/cases/{id}/restore')
def restore_case(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));x.archived=False;audit(s,u.id,'restore','test_case',id,x.project_id,after=data(x));s.commit();return {'ok':True}
@app.post('/api/cases/bulk-archive')
def bulk_archive_cases(project_id:int,case_ids:list[int],s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
    rows=s.scalars(select(TestCase).where(TestCase.project_id==project_id,TestCase.id.in_(case_ids))).all()
    if len(rows)!=len(set(case_ids)): raise HTTPException(400,'One or more cases do not belong to this project')
    for x in rows: x.archived=True;audit(s,u.id,'archive','test_case',x.id,project_id)
    s.commit();return {'archived':len(rows)}
@app.get('/api/folders')
def folders(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);return [data(x) for x in s.scalars(select(TestFolder).where(TestFolder.project_id==project_id).order_by(TestFolder.name)).all()]
@app.post('/api/folders')
def create_folder(project_id:int,name:str,parent_id:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
    if parent_id and obj(s,TestFolder,parent_id).project_id!=project_id: raise HTTPException(400,'Parent folder belongs to another project')
    if s.scalar(select(TestFolder).where(TestFolder.project_id==project_id,TestFolder.name==name)): raise HTTPException(409,'Folder name already exists')
    x=TestFolder(project_id=project_id,parent_id=parent_id,name=name);s.add(x);s.flush();audit(s,u.id,'create','folder',x.id,project_id,after=data(x));s.commit();return data(x)
@app.get('/api/custom-fields')
def custom_fields(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);return [data(x) for x in s.scalars(select(CustomField).where(CustomField.project_id==project_id)).all()]
@app.post('/api/custom-fields')
def create_custom_field(project_id:int,name:str,field_type:str,required:bool=False,options_json:str='',s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead'))
    if field_type not in ('text','number','date','boolean','select'): raise HTTPException(400,'Invalid field type')
    try: options=json.loads(options_json) if options_json else []
    except json.JSONDecodeError as exc: raise HTTPException(400,'Options must be JSON') from exc
    if not isinstance(options,list) or any(not isinstance(option,str) for option in options): raise HTTPException(400,'Options must be a list of strings')
    if field_type=='select' and not options: raise HTTPException(400,'Select fields need at least one option')
    if s.scalar(select(CustomField).where(CustomField.project_id==project_id,CustomField.name==name)): raise HTTPException(409,'Field name already exists')
    x=CustomField(project_id=project_id,name=name,field_type=field_type,required=required,options=options);s.add(x);s.flush();audit(s,u.id,'create','custom_field',x.id,project_id,after=data(x));s.commit();return data(x)
@app.patch('/api/custom-fields/{id}')
def update_custom_field(id:int,name:Optional[str]=None,required:Optional[bool]=None,options_json:Optional[str]=None,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,CustomField,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead'));before=data(x)
    if name is not None: x.name=name
    if required is not None: x.required=required
    if options_json is not None:
        try: options=json.loads(options_json)
        except json.JSONDecodeError as exc: raise HTTPException(400,'Options must be JSON') from exc
        if not isinstance(options,list) or any(not isinstance(item,str) for item in options): raise HTTPException(400,'Options must be a list of strings')
        if x.field_type=='select' and not options: raise HTTPException(400,'Select fields need at least one option')
        x.options=options
    audit(s,u.id,'update','custom_field',id,x.project_id,before,data(x));s.commit();return data(x)
@app.delete('/api/custom-fields/{id}')
def delete_custom_field(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,CustomField,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead'))
    for case_record in s.scalars(select(TestCase).where(TestCase.project_id==x.project_id)).all():
        values=dict(case_record.custom_values or {});values.pop(x.name,None);case_record.custom_values=values
    audit(s,u.id,'delete','custom_field',id,x.project_id,before=data(x));s.delete(x);s.commit();return {'ok':True}
@app.get('/api/plans')
def plans(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);return [data(x) for x in s.scalars(select(TestPlan).where(TestPlan.project_id==project_id,TestPlan.archived==False)).all()]
@app.post('/api/plans')
def create_plan(b:PlanIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead'))
    if b.owner_id: obj(s,User,b.owner_id)
    x=TestPlan(project_id=b.project_id,name=b.name,release=b.release,scope=b.scope,owner_id=b.owner_id);s.add(x);s.flush()
    for case_id in set(b.case_ids):
        c=obj(s,TestCase,case_id)
        if c.project_id!=b.project_id: raise HTTPException(400,'Plan case is from another project')
        s.add(PlanCase(plan_id=x.id,case_id=case_id))
    audit(s,u.id,'create','test_plan',x.id,b.project_id,after=data(x));s.commit();return data(x)
@app.post('/api/plans/{id}/cases')
def add_plan_case(id:int,case_ids:list[int],s:Session=Depends(db),u:User=Depends(user)):
    p=obj(s,TestPlan,id);accessible(s,u,p.project_id);require_project_role(s,u,p.project_id,('admin','lead'))
    for case_id in set(case_ids):
        c=obj(s,TestCase,case_id)
        if c.project_id!=p.project_id: raise HTTPException(400,'Case belongs to another project')
        if not s.scalar(select(PlanCase).where(PlanCase.plan_id==id,PlanCase.case_id==case_id)): s.add(PlanCase(plan_id=id,case_id=case_id))
    audit(s,u.id,'add_cases','test_plan',id,p.project_id,after={'case_ids':case_ids});s.commit();return {'ok':True}
@app.get('/api/plans/{id}/cases')
def plan_cases(id:int,s:Session=Depends(db),u:User=Depends(user)):
    p=obj(s,TestPlan,id);accessible(s,u,p.project_id);return [data(c) for c in s.scalars(select(TestCase).join(PlanCase).where(PlanCase.plan_id==id)).all()]
@app.get('/api/cycles')
def cycles(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);return [data(x) for x in s.scalars(select(Cycle).join(TestPlan,Cycle.plan_id==TestPlan.id).where(TestPlan.project_id==project_id)).all()]
@app.post('/api/cycles')
def create_cycle(b:CycleIn,s:Session=Depends(db),u:User=Depends(user)):
    p=obj(s,TestPlan,b.plan_id);accessible(s,u,p.project_id);require_project_role(s,u,p.project_id,('admin','lead'))
    if b.start_date and b.end_date and b.end_date<b.start_date: raise HTTPException(400,'End date must be after start date')
    x=Cycle(**b.model_dump());s.add(x);s.flush();audit(s,u.id,'create','cycle',x.id,p.project_id,after=data(x));s.commit();return data(x)
@app.get('/api/cycles/{id}/cases')
def cycle_cases(id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_for_cycle(s,id));result=[]
    for item in s.scalars(select(CycleCase).where(CycleCase.cycle_id==id)).all():
        case_record=obj(s,TestCase,item.case_id);version=obj(s,CaseVersion,item.version_id)
        latest=s.scalar(select(Execution).where(Execution.cycle_case_id==item.id).order_by(Execution.id.desc()).limit(1))
        result.append(data(item)|{'title':case_record.title,'version':version.version,'steps':version.steps,'preconditions':version.preconditions,'last_status':latest.status if latest else None,'latest_execution_id':latest.id if latest else None})
    return result
@app.post('/api/cycles/{id}/cases')
def assign_case(id:int,b:AssignIn,s:Session=Depends(db),u:User=Depends(user)):
    pid=project_for_cycle(s,id);accessible(s,u,pid);require_project_role(s,u,pid,('admin','lead'));x=obj(s,TestCase,b.case_id)
    if x.project_id!=pid: raise HTTPException(400,'Case is from another project')
    if b.assigned_to:
        assignee=obj(s,User,b.assigned_to)
        if not s.scalar(select(Membership).where(Membership.project_id==pid,Membership.user_id==assignee.id)): raise HTTPException(400,'Assignee is not a project member')
    existing=s.scalar(select(CycleCase).where(CycleCase.cycle_id==id,CycleCase.case_id==x.id))
    if existing: raise HTTPException(409,'Already in cycle')
    version=b.version_id or x.version;v=s.scalar(select(CaseVersion).where(CaseVersion.case_id==x.id,CaseVersion.version==version))
    if not v: raise HTTPException(400,'Requested case version does not exist')
    item=CycleCase(cycle_id=id,case_id=x.id,version_id=v.id,assigned_to=b.assigned_to);s.add(item)
    if b.assigned_to: notify(s,b.assigned_to,'Test assigned',f'Test case {x.title} was assigned in cycle #{id}.',f'/cycles/{id}')
    audit(s,u.id,'assign_case','cycle_case',x.id,pid,after={'cycle_id':id,'case_id':x.id,'version_id':v.id,'assigned_to':b.assigned_to});s.commit();return data(item)
@app.post('/api/cycle-cases/{id}/executions')
def execute(id:int,b:ExecuteIn,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,CycleCase,id);pid=project_for_cycle(s,x.cycle_id);accessible(s,u,pid);require_project_role(s,u,pid,('admin','lead','tester'))
    if project_role(s,u,pid)=='tester' and x.assigned_to!=u.id: raise HTTPException(403,'Testers may only execute tests assigned to them')
    if b.status not in ('Passed','Failed','Blocked','Skipped','Not Run'): raise HTTPException(400,'Invalid status')
    if b.status=='Not Run': raise HTTPException(400,'Not Run is a cycle state, not an execution attempt')
    snap=obj(s,CaseVersion,x.version_id)
    valid_steps=set(range(len(snap.steps or [])))
    for result in b.step_results:
        if result.get('step') not in valid_steps or result.get('status') not in ('Passed','Failed','Blocked','Skipped','Not Run'):
            raise HTTPException(400,'Each step result needs a valid step index and status')
    e=Execution(cycle_case_id=id,user_id=u.id,status=b.status,notes=b.notes,actual_result=b.actual_result,step_results=b.step_results);s.add(e);s.flush();audit(s,u.id,'execute','execution',e.id,pid,after=data(e));s.commit();return data(e)
@app.get('/api/executions')
def executions(cycle_id:int,s:Session=Depends(db),u:User=Depends(user),offset:int=Query(0,ge=0),limit:int=Query(100,ge=1,le=500)):
    accessible(s,u,project_for_cycle(s,cycle_id));return [data(x) for x in s.scalars(select(Execution).join(CycleCase).where(CycleCase.cycle_id==cycle_id).order_by(Execution.id.desc()).offset(offset).limit(limit)).all()]
@app.get('/api/defects')
def defects(project_id:int,q:str='',status:Optional[str]=None,s:Session=Depends(db),u:User=Depends(user),offset:int=Query(0,ge=0),limit:int=Query(100,ge=1,le=500)):
    accessible(s,u,project_id);query=select(Defect).where(Defect.project_id==project_id)
    if q: query=query.where(Defect.title.ilike(f'%{q}%'))
    if status: query=query.where(Defect.status==status)
    return [data(x) for x in s.scalars(query.order_by(Defect.id.desc()).offset(offset).limit(limit)).all()]
@app.post('/api/defects')
def create_defect(b:DefectIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead','tester'))
    if b.severity not in ('Low','Medium','High','Critical') or b.priority not in ('Low','Medium','High','Critical'): raise HTTPException(400,'Invalid defect severity or priority')
    execution_ids=set(b.execution_ids)
    if b.execution_id: execution_ids.add(b.execution_id)
    for execution_id in execution_ids:
        e=obj(s,Execution,execution_id);cc=obj(s,CycleCase,e.cycle_case_id)
        if project_for_cycle(s,cc.cycle_id)!=b.project_id: raise HTTPException(400,'Execution belongs to another project')
    if b.assigned_to and not s.scalar(select(Membership).where(Membership.project_id==b.project_id,Membership.user_id==b.assigned_to)): raise HTTPException(400,'Assignee is not a project member')
    x=Defect(project_id=b.project_id,title=b.title,description=b.description,severity=b.severity,priority=b.priority,assigned_to=b.assigned_to);s.add(x);s.flush()
    for execution_id in execution_ids: s.add(DefectExecution(defect_id=x.id,execution_id=execution_id))
    audit(s,u.id,'create','defect',x.id,b.project_id,after=data(x));s.commit();return data(x)
@app.get('/api/defects/{id}')
def defect(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,Defect,id);accessible(s,u,x.project_id);return data(x)|{'execution_ids':[rel.execution_id for rel in s.scalars(select(DefectExecution).where(DefectExecution.defect_id==id)).all()]}
@app.patch('/api/defects/{id}')
def edit_defect(id:int,title:Optional[str]=None,description:Optional[str]=None,status:Optional[str]=None,severity:Optional[str]=None,priority:Optional[str]=None,assigned_to:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,Defect,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));before=data(x)
    if status is not None and status not in ('New','Open','In Progress','Fixed','Retest','Closed','Reopened'): raise HTTPException(400,'Invalid status')
    if severity is not None and severity not in ('Low','Medium','High','Critical'): raise HTTPException(400,'Invalid severity')
    if assigned_to and not s.scalar(select(Membership).where(Membership.project_id==x.project_id,Membership.user_id==assigned_to)): raise HTTPException(400,'Assignee is not a project member')
    for key,value in {'title':title,'description':description,'status':status,'severity':severity,'priority':priority,'assigned_to':assigned_to}.items():
        if value is not None: setattr(x,key,value)
    x.updated_at=now()
    if x.assigned_to and (x.assigned_to!=u.id or x.assigned_to!=before['assigned_to']): notify(s,x.assigned_to,'Defect updated',f'Defect {x.title}: {x.status}.',f'/defects/{id}')
    audit(s,u.id,'update','defect',id,x.project_id,before,data(x));s.commit();return data(x)
@app.post('/api/defects/{id}/executions')
def link_defect_execution(id:int,execution_id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,Defect,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'))
    e=obj(s,Execution,execution_id);cc=obj(s,CycleCase,e.cycle_case_id)
    if project_for_cycle(s,cc.cycle_id)!=x.project_id: raise HTTPException(400,'Execution belongs to another project')
    if not s.scalar(select(DefectExecution).where(DefectExecution.defect_id==id,DefectExecution.execution_id==execution_id)): s.add(DefectExecution(defect_id=id,execution_id=execution_id))
    s.commit();return {'ok':True}
@app.post('/api/automation/mappings')
def create_mapping(b:MapIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead'));x=obj(s,TestCase,b.case_id)
    if x.project_id!=b.project_id: raise HTTPException(400,'Wrong project')
    m=AutomationMap(**b.model_dump());s.add(m)
    try:s.commit()
    except Exception:s.rollback();raise HTTPException(409,'Mapping already exists')
    return data(m)
@app.get('/api/automation/mappings')
def mappings(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);return [data(x) for x in s.scalars(select(AutomationMap).where(AutomationMap.project_id==project_id)).all()]
@app.post('/api/automation/import/{cycle_id}')
async def import_junit(cycle_id:int,file:UploadFile=File(...),s:Session=Depends(db),u:User=Depends(user)):
    pid=project_for_cycle(s,cycle_id);accessible(s,u,pid);require_project_role(s,u,pid,('admin','lead'))
    if not (file.filename or '').lower().endswith('.xml'): raise HTTPException(415,'Upload a .xml JUnit report')
    raw=await file.read(5_000_001)
    if len(raw)>5_000_000: raise HTTPException(413,'XML file too large')
    return import_junit_content(cycle_id,file.filename or '',raw,s,u,pid)

def import_junit_content(cycle_id,filename,raw,s,u,pid):
    checksum=_hashlib.sha256(raw).hexdigest()
    prior=s.scalar(select(ImportLog).where(ImportLog.cycle_id==cycle_id,ImportLog.checksum==checksum))
    if prior: raise HTTPException(409,f'Duplicate import (import #{prior.id})')
    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper(): raise HTTPException(400,'DTD/entities not allowed')
    try:root=ET.fromstring(raw)
    except ET.ParseError as exc: raise HTTPException(400,'Invalid XML') from exc
    cases=root.findall('.//testcase')
    if not cases:
        cases=root.findall('.//{*}testcase')
    imported=0;unmatched=[];errors=[]
    for tc in cases:
        key=(tc.get('classname','')+'.'+tc.get('name','')).strip('.')
        mapping=s.scalar(select(AutomationMap).where(AutomationMap.project_id==pid,AutomationMap.external_id==key))
        if not mapping: unmatched.append(key);continue
        cc=s.scalar(select(CycleCase).where(CycleCase.cycle_id==cycle_id,CycleCase.case_id==mapping.case_id))
        if not cc: unmatched.append(key+' (not in cycle)');continue
        failure=tc.find('failure')
        if failure is None: failure=tc.find('error')
        if failure is None: failure=tc.find('{*}failure')
        if failure is None: failure=tc.find('{*}error')
        skipped=tc.find('skipped')
        if skipped is None: skipped=tc.find('{*}skipped')
        status='Failed' if failure is not None else 'Skipped' if skipped is not None else 'Passed'
        notes=((failure.text or failure.get('message',''))[:4000] if failure is not None else '')
        execution=Execution(cycle_case_id=cc.id,user_id=u.id,status=status,source='Automation',notes=notes,duration=tc.get('time',''));s.add(execution);imported+=1
    log=ImportLog(cycle_id=cycle_id,filename=os.path.basename(filename)[:255],checksum=checksum,imported=imported,unmatched=unmatched,errors=errors);s.add(log);s.flush();audit(s,u.id,'import','junit',log.id,pid,after={'filename':log.filename,'checksum':checksum,'imported':imported,'unmatched':unmatched});s.commit()
    return {'import_id':log.id,'imported':imported,'unmatched':unmatched,'errors':errors,'duplicate':False}
@app.get('/api/automation/imports')
def automation_imports(cycle_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_for_cycle(s,cycle_id))
    return [data(x) for x in s.scalars(select(ImportLog).where(ImportLog.cycle_id==cycle_id).order_by(ImportLog.id.desc())).all()]
@app.post('/api/automation/tokens')
def create_api_token(project_id:int,name:str,expires_at:Optional[datetime]=None,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead'))
    secret=secrets.token_urlsafe(32);x=ApiToken(name=name,token_hash=_hashlib.sha256(secret.encode()).hexdigest(),prefix=secret[:10],user_id=u.id,project_id=project_id,expires_at=expires_at)
    s.add(x);s.flush();audit(s,u.id,'create','api_token',x.id,project_id,after={'name':name,'prefix':x.prefix,'expires_at':str(expires_at)});s.commit()
    return {'id':x.id,'name':name,'token':secret,'prefix':x.prefix,'expires_at':expires_at}
@app.get('/api/automation/tokens')
def list_api_tokens(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead'))
    return [{'id':x.id,'name':x.name,'prefix':x.prefix,'expires_at':x.expires_at,'revoked':x.revoked,'created_at':x.created_at} for x in s.scalars(select(ApiToken).where(ApiToken.project_id==project_id)).all()]
@app.delete('/api/automation/tokens/{id}')
def revoke_api_token(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,ApiToken,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead'));x.revoked=True;audit(s,u.id,'revoke','api_token',id,x.project_id);s.commit();return {'ok':True}
@app.post('/api/automation/ci/{cycle_id}')
async def automation_ci_import(cycle_id:int,file:UploadFile=File(...),credentials:HTTPAuthorizationCredentials=Depends(HTTPBearer()),s:Session=Depends(db)):
    token_hash=_hashlib.sha256(credentials.credentials.encode()).hexdigest()
    api_token=s.scalar(select(ApiToken).where(ApiToken.token_hash==token_hash,ApiToken.revoked==False))
    expiry=api_token.expires_at.replace(tzinfo=timezone.utc) if api_token and api_token.expires_at and api_token.expires_at.tzinfo is None else (api_token.expires_at if api_token else None)
    if not api_token or (expiry and expiry<=now()): raise HTTPException(401,'Invalid or expired integration token')
    pid=project_for_cycle(s,cycle_id)
    if api_token.project_id!=pid: raise HTTPException(403,'Integration token is not authorized for this project')
    actor=obj(s,User,api_token.user_id)
    raw=await file.read(5_000_001)
    if len(raw)>5_000_000: raise HTTPException(413,'XML file too large')
    return import_junit_content(cycle_id,file.filename or '',raw,s,actor,pid)
@app.post('/api/executions/{id}/attachments')
async def upload_attachment(id:int,file:UploadFile=File(...),s:Session=Depends(db),u:User=Depends(user)):
    e=obj(s,Execution,id);cc=obj(s,CycleCase,e.cycle_case_id);pid=project_for_cycle(s,cc.cycle_id);accessible(s,u,pid)
    if file.size and file.size>MAX_UPLOAD_BYTES: raise HTTPException(413,'Attachment exceeds the size limit')
    raw=await file.read(MAX_UPLOAD_BYTES+1)
    if len(raw)>MAX_UPLOAD_BYTES: raise HTTPException(413,'Attachment exceeds the size limit')
    original=os.path.basename(file.filename or '').replace('\x00','')
    ext=os.path.splitext(original)[1].lower()
    allowed={'.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg','.webp':'image/webp','.txt':'text/plain','.log':'text/plain'}
    if ext not in allowed: raise HTTPException(415,'Only PNG, JPEG, WEBP, TXT and LOG files are allowed')
    mime=allowed[ext]
    signatures={'.png':raw.startswith(b'\x89PNG\r\n\x1a\n'),'.jpg':raw.startswith(b'\xff\xd8\xff'),'.jpeg':raw.startswith(b'\xff\xd8\xff'),'.webp':len(raw)>=12 and raw[:4]==b'RIFF' and raw[8:12]==b'WEBP'}
    if ext in signatures and not signatures[ext]: raise HTTPException(400,'Image content does not match its file extension')
    if mime=='text/plain' and b'\x00' in raw: raise HTTPException(400,'Text files must not contain binary data')
    key=secrets.token_hex(24);path=os.path.join(UPLOAD_DIR,key)
    os.makedirs(UPLOAD_DIR,exist_ok=True)
    with open(path,'xb') as stream: stream.write(raw)
    x=Attachment(project_id=pid,execution_id=id,filename=original[:255] or 'attachment'+ext,content_type=mime,size=len(raw),storage_key=key,uploaded_by=u.id);s.add(x);s.flush();audit(s,u.id,'upload','attachment',x.id,pid,after={'filename':x.filename,'size':x.size});s.commit();return data(x)
@app.get('/api/attachments/{id}')
def get_attachment(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,Attachment,id);accessible(s,u,x.project_id);path=os.path.join(UPLOAD_DIR,x.storage_key)
    if not os.path.isfile(path): raise HTTPException(404,'Attachment file not found')
    return FileResponse(path,media_type=x.content_type,filename=x.filename,content_disposition_type='attachment')
@app.get('/api/executions/{id}/attachments')
def execution_attachments(id:int,s:Session=Depends(db),u:User=Depends(user)):
    e=obj(s,Execution,id);cc=obj(s,CycleCase,e.cycle_case_id);accessible(s,u,project_for_cycle(s,cc.cycle_id))
    return [data(x) for x in s.scalars(select(Attachment).where(Attachment.execution_id==id)).all()]
@app.post('/api/cases/import.xlsx')
async def import_cases_xlsx(project_id:int,file:UploadFile=File(...),s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
    if not (file.filename or '').lower().endswith('.xlsx'): raise HTTPException(415,'Upload an .xlsx workbook')
    raw=await file.read(5_000_001)
    if len(raw)>5_000_000: raise HTTPException(413,'Workbook exceeds 5 MB')
    try:
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    except Exception as exc: raise HTTPException(400,'Invalid Excel workbook') from exc
    sheet=workbook.active
    rows=sheet.iter_rows(values_only=True)
    headers=[str(v or '').strip().lower() for v in next(rows,())]
    required={'title','priority','kind'}
    if not required.issubset(headers): raise HTTPException(400,'Workbook needs title, priority and kind columns')
    index={name:headers.index(name) for name in headers}
    imported=0;errors=[]
    for row_number,row in enumerate(rows,start=2):
        title=str(row[index['title']] or '').strip() if index['title']<len(row) else ''
        if not title:
            errors.append({'row':row_number,'error':'Title is required'});continue
        def cell(name,default=''):
            i=index.get(name)
            return str(row[i] if i is not None and i<len(row) and row[i] is not None else default)
        try: steps=json.loads(cell('steps','[]'))
        except json.JSONDecodeError:
            errors.append({'row':row_number,'error':'Steps must contain JSON'});continue
        if not isinstance(steps,list):
            errors.append({'row':row_number,'error':'Steps must be a JSON array'});continue
        case=TestCase(project_id=project_id,title=title[:300],priority=cell('priority','Medium'),kind=cell('kind','Manual'),status=cell('status','Active'),author_id=u.id)
        s.add(case);s.flush();s.add(CaseVersion(case_id=case.id,version=1,preconditions=cell('preconditions'),steps=steps,author_id=u.id));imported+=1
    audit(s,u.id,'import','test_cases',project_id,project_id,after={'filename':os.path.basename(file.filename or ''),'imported':imported,'errors':errors});s.commit();return {'imported':imported,'errors':errors}
@app.get('/api/cases/export/xlsx')
def export_cases_xlsx(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id)
    from openpyxl import Workbook
    workbook=Workbook();sheet=workbook.active;sheet.title='Test Cases'
    sheet.append(['key','title','priority','kind','status','preconditions','steps','version'])
    for c in s.scalars(select(TestCase).where(TestCase.project_id==project_id).order_by(TestCase.id)).all():
        v=s.scalar(select(CaseVersion).where(CaseVersion.case_id==c.id,CaseVersion.version==c.version))
        sheet.append([f'TC-{c.id}',c.title,c.priority,c.kind,c.status,v.preconditions if v else '',json.dumps(v.steps if v else []),c.version])
    buf=io.BytesIO();workbook.save(buf);buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="test-cases.xlsx"'})
@app.get('/api/notifications')
def notifications(s:Session=Depends(db),u:User=Depends(user)):
    return [data(x) for x in s.scalars(select(Notification).where(Notification.user_id==u.id).order_by(Notification.created_at.desc()).limit(100)).all()]
@app.post('/api/notifications/{id}/read')
def mark_notification_read(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,Notification,id)
    if x.user_id!=u.id: raise HTTPException(404,'Notification not found')
    x.read=True;s.commit();return {'ok':True}
@app.post('/api/requirements')
def create_requirement(b:RequirementIn,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead','tester'))
    x=RequirementRecord(**b.model_dump());s.add(x);s.flush();audit(s,u.id,'create','requirement',x.id,b.project_id,after=data(x));s.commit();return data(x)
@app.patch('/api/requirements/{id}')
def update_requirement(id:int,b:RequirementIn,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,RequirementRecord,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'))
    if b.project_id!=x.project_id: raise HTTPException(400,'Cannot move a requirement across projects')
    before=data(x);x.key=b.key;x.title=b.title;x.description=b.description;x.status=b.status;audit(s,u.id,'update','requirement',id,x.project_id,before,data(x));s.commit();return data(x)
@app.delete('/api/requirements/{id}')
def delete_requirement(id:int,s:Session=Depends(db),u:User=Depends(user)):
    x=obj(s,RequirementRecord,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));audit(s,u.id,'delete','requirement',id,x.project_id,before=data(x))
    for link in s.scalars(select(RequirementCase).where(RequirementCase.requirement_id==id)).all(): s.delete(link)
    s.delete(x);s.commit();return {'ok':True}
@app.get('/api/requirements')
def requirements(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id)
    return [data(r)|{'case_ids':[x.case_id for x in s.scalars(select(RequirementCase).where(RequirementCase.requirement_id==r.id)).all()]} for r in s.scalars(select(RequirementRecord).where(RequirementRecord.project_id==project_id)).all()]
@app.post('/api/requirements/{id}/cases')
def link_requirement_case(id:int,case_id:int,s:Session=Depends(db),u:User=Depends(user)):
    req=obj(s,RequirementRecord,id);accessible(s,u,req.project_id);require_project_role(s,u,req.project_id,('admin','lead','tester'));c=obj(s,TestCase,case_id)
    if c.project_id!=req.project_id: raise HTTPException(400,'Case belongs to another project')
    if not s.scalar(select(RequirementCase).where(RequirementCase.requirement_id==id,RequirementCase.case_id==case_id)): s.add(RequirementCase(requirement_id=id,case_id=case_id))
    s.commit();return {'ok':True}
@app.delete('/api/requirements/{id}/cases/{case_id}')
def unlink_requirement_case(id:int,case_id:int,s:Session=Depends(db),u:User=Depends(user)):
    req=obj(s,RequirementRecord,id);accessible(s,u,req.project_id);require_project_role(s,u,req.project_id,('admin','lead','tester'))
    link=s.scalar(select(RequirementCase).where(RequirementCase.requirement_id==id,RequirementCase.case_id==case_id))
    if link: s.delete(link);s.commit()
    return {'ok':True}
@app.get('/api/requirements/coverage')
def requirement_coverage(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);result=[]
    for r in s.scalars(select(RequirementRecord).where(RequirementRecord.project_id==project_id)).all():
        links=s.scalars(select(RequirementCase).where(RequirementCase.requirement_id==r.id)).all()
        result.append(data(r)|{'case_ids':[x.case_id for x in links],'covered':bool(links)})
    return result
@app.get('/api/audit')
def audit_events(project_id:Optional[int]=None,limit:int=Query(100,ge=1,le=500),s:Session=Depends(db),u:User=Depends(user)):
    query=select(AuditEvent).order_by(AuditEvent.created_at.desc()).limit(limit)
    if project_id is not None:
        accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead'));query=query.where(AuditEvent.project_id==project_id)
    elif not has_any_role(u,'admin'): raise HTTPException(403,'Admin required')
    return [data(x) for x in s.scalars(query).all()]
@app.get('/api/dashboard')
def dashboard(project_id:int,cycle_id:Optional[int]=None,release:Optional[str]=None,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id)
    case_count=s.scalar(select(func.count()).select_from(TestCase).where(TestCase.project_id==project_id,TestCase.archived==False))
    defects=s.scalars(select(Defect).where(Defect.project_id==project_id)).all()
    cycle_query=select(Cycle.id).join(TestPlan).where(TestPlan.project_id==project_id)
    if release is not None: cycle_query=cycle_query.where(TestPlan.release==release)
    if cycle_id is not None:
        if project_for_cycle(s,cycle_id)!=project_id: raise HTTPException(400,'Cycle belongs to another project')
        cycle_query=cycle_query.where(Cycle.id==cycle_id)
    cycle_ids=s.scalars(cycle_query).all()
    entries=s.scalars(select(Execution).join(CycleCase).where(CycleCase.cycle_id.in_(cycle_ids))).all() if cycle_ids else []
    cycle_items=s.scalars(select(CycleCase).where(CycleCase.cycle_id.in_(cycle_ids))).all() if cycle_ids else []
    latest={item.id:s.scalar(select(Execution.status).where(Execution.cycle_case_id==item.id).order_by(Execution.id.desc()).limit(1)) for item in cycle_items}
    statuses={k:sum(status==k for status in latest.values()) for k in ('Passed','Failed','Blocked','Skipped')}
    statuses['Not Run']=sum(status is None for status in latest.values())
    completed=sum(status is not None for status in latest.values())
    recent=s.scalars(select(Execution).join(CycleCase).where(CycleCase.cycle_id.in_(cycle_ids)).order_by(Execution.created_at.desc()).limit(10)).all() if cycle_ids else []
    return {'cases':case_count,'executions':len(entries),'statuses':statuses,'open_defects':sum(d.status!='Closed' for d in defects),'defects_by_severity':{k:sum(d.status!='Closed' and d.severity==k for d in defects) for k in ('Critical','High','Medium','Low')},'cycles':len(cycle_ids),'progress':(completed/len(cycle_items) if cycle_items else 0),'recent_activity':[data(e) for e in recent],'recent_defects':[data(d) for d in defects[-5:]]}
@app.get('/api/reports/executions.csv')
def report(project_id:int,cycle_id:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);buf=io.StringIO();w=csv.writer(buf);w.writerow(['execution_id','case_id','cycle_id','status','source','user_id','timestamp'])
    query=select(Execution,CycleCase).join(CycleCase,Execution.cycle_case_id==CycleCase.id).join(Cycle,CycleCase.cycle_id==Cycle.id).join(TestPlan,Cycle.plan_id==TestPlan.id).where(TestPlan.project_id==project_id)
    if cycle_id is not None:
        if project_for_cycle(s,cycle_id)!=project_id: raise HTTPException(400,'Cycle belongs to another project')
        query=query.where(CycleCase.cycle_id==cycle_id)
    rows=s.execute(query).all()
    for e,cc in rows:w.writerow([e.id,cc.case_id,cc.cycle_id,e.status,e.source,e.user_id,e.created_at])
    buf.seek(0);return StreamingResponse(iter([buf.getvalue()]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="executions.csv"'})
@app.get('/api/reports/defects.csv')
def defect_report(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
    accessible(s,u,project_id);buf=io.StringIO();w=csv.writer(buf);w.writerow(['defect_id','title','severity','priority','status','assigned_to','created_at'])
    for d in s.scalars(select(Defect).where(Defect.project_id==project_id).order_by(Defect.id)).all(): w.writerow([d.id,d.title,d.severity,d.priority,d.status,d.assigned_to,d.created_at])
    buf.seek(0);return StreamingResponse(iter([buf.getvalue()]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="defects.csv"'})
@app.get('/api/health')
def health(s:Session=Depends(db)):
    s.execute(select(1))
    return {'status':'ok','database':'ok','api_version':'v1'}
