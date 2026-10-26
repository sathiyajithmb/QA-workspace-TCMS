"""Test case repository, revisions, folders, and Excel import/export routes."""
import io
import json
import os
import re
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import User, TestCase, CaseVersion, TestFolder, now

class CaseIn(BaseModel): project_id:int; title:str; priority:str='Medium'; kind:str='Manual'; status:str='Active'; preconditions:str=''; steps:list[dict]=Field(default_factory=list); folder_id:Optional[int]=None; custom_values:dict=Field(default_factory=dict); import_metadata:Optional[dict]=None

def create_test_cases_router(db, user, accessible, require_project_role, obj, data, audit, validate_custom_values):
    router = APIRouter(tags=["Test Cases"])

    @router.post('/api/cases/import.xlsx')
    async def import_cases_xlsx(project_id:int,file:UploadFile=File(...),s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
        if not (file.filename or '').lower().endswith('.xlsx'): raise HTTPException(415,'Upload an .xlsx workbook')
        raw=await file.read(5_000_001)
        if len(raw)>5_000_000: raise HTTPException(413,'Workbook exceeds 5 MB')
        try:
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        except Exception as exc: raise HTTPException(400,'Invalid Excel workbook') from exc
        try:
            sheet=workbook.active
            rows=sheet.iter_rows(values_only=True)
            headers=[re.sub(r'[_\s]+', ' ', str(value or '').strip().lower()) for value in next(rows,())]
            sheet_format='test scenario' in headers
            required={'test scenario','priority','test type'} if sheet_format else {'title','priority','kind'}
            if not required.issubset(headers):
                raise HTTPException(400,'Workbook needs Test Scenario, Priority and Test Type columns (or title, priority and kind)')
            named=[name for name in headers if name]
            if len(named)!=len(set(named)):
                raise HTTPException(400,'Duplicate column headers are not supported')
            index={name:i for i,name in enumerate(headers)}
            imported=0;errors=[];case_ids=[]
            for row_number,row in enumerate(rows,start=2):
                if all(value is None or str(value).strip()=='' for value in row):
                    continue
                def cell(*names, default=''):
                    for name in names:
                        position=index.get(name)
                        if position is not None and position<len(row) and row[position] is not None:
                            return str(row[position]).strip()
                    return default
                title=cell('test scenario' if sheet_format else 'title')
                if not title:
                    errors.append({'row':row_number,'error':'Title is required'});continue
                if len(title)>300:
                    errors.append({'row':row_number,'error':'Test scenario must be at most 300 characters'});continue
                priority=(cell('priority') or 'Medium').title()
                kind=(cell('kind') or 'Manual').title()
                source_status=cell('status')
                status='Active' if sheet_format else (source_status or 'Active').title()
                if priority not in ('Low','Medium','High','Critical') or kind not in ('Manual','Automated','Both') or status not in ('Active','Draft','Deprecated'):
                    errors.append({'row':row_number,'error':'Invalid priority, kind or status'});continue
                expected=cell('expected result','expected results')
                raw_steps=cell('test steps' if sheet_format else 'steps')
                if sheet_format:
                    # The sheet has one overall expected result, attached to the final step.
                    actions=[re.sub(r'^\s*\d+[.)]\s*', '', line).strip() for line in raw_steps.splitlines() if line.strip()]
                    steps=[{'action':action,'expected':expected if i==len(actions)-1 else ''} for i,action in enumerate(actions)]
                else:
                    try: steps=json.loads(raw_steps or '[]')
                    except json.JSONDecodeError:
                        errors.append({'row':row_number,'error':'Steps must contain JSON'});continue
                if not isinstance(steps,list):
                    errors.append({'row':row_number,'error':'Steps must be a JSON array'});continue
                if any(not isinstance(step,dict) for step in steps):
                    errors.append({'row':row_number,'error':'Each step must be a JSON object'});continue
                metadata={
                    'source_id':cell('source id','test case id','testcase id','case id','-',''),
                    'feature':cell('feature'),
                    'test_type':cell('test type'),
                    'actual_result':cell('actual result'),
                    'sheet_status':source_status if sheet_format else cell('sheet status'),
                    'expected_result':expected,
                    'test_steps':raw_steps if sheet_format else cell('test steps'),
                }
                metadata={key:value for key,value in metadata.items() if value}
                case=TestCase(project_id=project_id,title=title,priority=priority,kind=kind,status=status,author_id=u.id,import_metadata=metadata)
                s.add(case);s.flush()
                s.add(CaseVersion(case_id=case.id,version=1,preconditions=cell('preconditions','pre conditions'),steps=steps,author_id=u.id))
                imported+=1;case_ids.append(case.id)
            audit(s,u.id,'import','test_cases',project_id,project_id,after={'filename':os.path.basename(file.filename or ''),'imported':imported,'errors':errors});s.commit();return {'imported':imported,'errors':errors,'case_ids':case_ids}
        finally:
            workbook.close()
    @router.get('/api/cases/export/xlsx')
    def export_cases_xlsx(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id)
        from openpyxl import Workbook
        workbook=Workbook();sheet=workbook.active;sheet.title='Test Cases'
        sheet.append(['key','title','priority','kind','status','preconditions','steps','version','source_id','feature','test_type','actual_result','sheet_status','expected_result','test_steps'])
        for c in s.scalars(select(TestCase).where(TestCase.project_id==project_id).order_by(TestCase.id)).all():
            v=s.scalar(select(CaseVersion).where(CaseVersion.case_id==c.id,CaseVersion.version==c.version))
            sheet.append([f'TC-{c.id}',c.title,c.priority,c.kind,c.status,v.preconditions if v else '',json.dumps(v.steps if v else []),c.version]+[(c.import_metadata or {}).get(key,'') for key in ('source_id','feature','test_type','actual_result','sheet_status','expected_result','test_steps')])
        buf=io.BytesIO();workbook.save(buf);buf.seek(0)
        return StreamingResponse(iter([buf.getvalue()]),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="test-cases.xlsx"'})
    @router.get('/api/cases')
    def cases(project_id:int,q:str='',priority:Optional[str]=None,kind:Optional[str]=None,folder_id:Optional[int]=None,include_archived:bool=False,offset:int=Query(0,ge=0),limit:int=Query(100,ge=1,le=500),s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id)
        query=select(TestCase,CaseVersion).outerjoin(CaseVersion,(CaseVersion.case_id==TestCase.id)&(CaseVersion.version==TestCase.version)).where(TestCase.project_id==project_id)
        if not include_archived: query=query.where(TestCase.archived==False)
        if q: query=query.where(TestCase.title.ilike(f'%{q}%'))
        if priority: query=query.where(TestCase.priority==priority)
        if kind: query=query.where(TestCase.kind==kind)
        if folder_id is not None: query=query.where(TestCase.folder_id==folder_id)
        return [data(case)|{'preconditions':version.preconditions if version else '', 'steps':version.steps if version else []} for case,version in s.execute(query.order_by(TestCase.id).offset(offset).limit(limit)).all()]
    @router.get('/api/cases/{id}')
    def case(id:int,version:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
        x=obj(s,TestCase,id);accessible(s,u,x.project_id);v=s.scalar(select(CaseVersion).where(CaseVersion.case_id==id,CaseVersion.version==(version if version is not None else x.version)))
        if not v: raise HTTPException(404,'Test case version not found')
        return data(x)|{'preconditions':v.preconditions,'steps':v.steps,'version_number':v.version}
    @router.get('/api/cases/{id}/versions')
    def case_versions(id:int,s:Session=Depends(db),u:User=Depends(user)):
        x=obj(s,TestCase,id);accessible(s,u,x.project_id)
        return [data(v) for v in s.scalars(select(CaseVersion).where(CaseVersion.case_id==id).order_by(CaseVersion.version.desc())).all()]
    @router.post('/api/cases')
    def create_case(b:CaseIn,s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,b.project_id);require_project_role(s,u,b.project_id,('admin','lead','tester'))
        validate_custom_values(s,b.project_id,b.custom_values)
        if b.priority not in ('Low','Medium','High','Critical') or b.kind not in ('Manual','Automated','Both') or b.status not in ('Active','Draft','Deprecated'): raise HTTPException(400,'Invalid test case metadata')
        if b.folder_id: 
            folder=obj(s,TestFolder,b.folder_id)
            if folder.project_id!=b.project_id: raise HTTPException(400,'Folder belongs to another project')
        x=TestCase(project_id=b.project_id,folder_id=b.folder_id,title=b.title,priority=b.priority,kind=b.kind,status=b.status,author_id=u.id,custom_values=b.custom_values,import_metadata=b.import_metadata or {});s.add(x);s.flush();s.add(CaseVersion(case_id=x.id,version=1,preconditions=b.preconditions,steps=b.steps,author_id=u.id));audit(s,u.id,'create','test_case',x.id,b.project_id,after=data(x));s.commit();return data(x)
    @router.put('/api/cases/{id}')
    def edit_case(id:int,b:CaseIn,s:Session=Depends(db),u:User=Depends(user)):
        x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));before=data(x)
        if b.project_id!=x.project_id: raise HTTPException(400,'Cannot move case across projects')
        validate_custom_values(s,x.project_id,b.custom_values)
        if b.priority not in ('Low','Medium','High','Critical') or b.kind not in ('Manual','Automated','Both') or b.status not in ('Active','Draft','Deprecated'): raise HTTPException(400,'Invalid test case metadata')
        if b.folder_id:
            folder=obj(s,TestFolder,b.folder_id)
            if folder.project_id!=x.project_id: raise HTTPException(400,'Folder belongs to another project')
        x.title=b.title;x.priority=b.priority;x.kind=b.kind;x.status=b.status;x.folder_id=b.folder_id;x.custom_values=b.custom_values
        if b.import_metadata is not None: x.import_metadata=b.import_metadata
        x.version+=1;x.updated_at=now();s.add(CaseVersion(case_id=id,version=x.version,preconditions=b.preconditions,steps=b.steps,author_id=u.id));audit(s,u.id,'update','test_case',id,x.project_id,before,data(x));s.commit();return data(x)
    @router.post('/api/cases/{id}/archive')
    def archive_case(id:int,s:Session=Depends(db),u:User=Depends(user)):
        x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));x.archived=True;audit(s,u.id,'archive','test_case',id,x.project_id,after=data(x));s.commit();return {'ok':True}
    @router.post('/api/cases/{id}/restore')
    def restore_case(id:int,s:Session=Depends(db),u:User=Depends(user)):
        x=obj(s,TestCase,id);accessible(s,u,x.project_id);require_project_role(s,u,x.project_id,('admin','lead','tester'));x.archived=False;audit(s,u.id,'restore','test_case',id,x.project_id,after=data(x));s.commit();return {'ok':True}
    @router.post('/api/cases/bulk-archive')
    def bulk_archive_cases(project_id:int,case_ids:list[int],s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
        rows=s.scalars(select(TestCase).where(TestCase.project_id==project_id,TestCase.id.in_(case_ids))).all()
        if len(rows)!=len(set(case_ids)): raise HTTPException(400,'One or more cases do not belong to this project')
        for x in rows: x.archived=True;audit(s,u.id,'archive','test_case',x.id,project_id)
        s.commit();return {'archived':len(rows)}
    @router.get('/api/folders')
    def folders(project_id:int,s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id);return [data(x) for x in s.scalars(select(TestFolder).where(TestFolder.project_id==project_id).order_by(TestFolder.name)).all()]
    @router.post('/api/folders')
    def create_folder(project_id:int,name:str,parent_id:Optional[int]=None,s:Session=Depends(db),u:User=Depends(user)):
        accessible(s,u,project_id);require_project_role(s,u,project_id,('admin','lead','tester'))
        if parent_id and obj(s,TestFolder,parent_id).project_id!=project_id: raise HTTPException(400,'Parent folder belongs to another project')
        if s.scalar(select(TestFolder).where(TestFolder.project_id==project_id,TestFolder.name==name)): raise HTTPException(409,'Folder name already exists')
        x=TestFolder(project_id=project_id,parent_id=parent_id,name=name);s.add(x);s.flush();audit(s,u.id,'create','folder',x.id,project_id,after=data(x));s.commit();return data(x)

    return router
