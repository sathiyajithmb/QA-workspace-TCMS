import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import main


@pytest.fixture
def client(tmp_path, monkeypatch):
    database = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{database}", connect_args={"check_same_thread": False})
    monkeypatch.setattr(main, "DATABASE_URL", f"sqlite:///{database}")
    monkeypatch.setattr(main, "engine", engine)
    monkeypatch.setattr(main, "SessionLocal", sessionmaker(bind=engine))
    monkeypatch.setattr(main, "UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(main, "SECRET", "test-only-session-secret")
    monkeypatch.setenv("ADMIN_EMAIL", "admin@example.test")
    monkeypatch.setenv("ADMIN_PASSWORD", "test-admin-password-2026")

    with TestClient(main.app) as test_client:
        response = test_client.post(
            "/api/auth/login",
            json={"email": "admin@example.test", "password": "test-admin-password-2026"},
        )
        assert response.status_code == 200
        assert "httponly" in response.headers["set-cookie"].lower()
        yield test_client
    engine.dispose()


def make_project(client):
    response = client.post(
        "/api/projects",
        json={"key": "DEMO", "name": "Demo Project", "description": "Acceptance tests"},
    )
    assert response.status_code == 200, response.text
    return response.json()


def make_case(client, project_id, title="Login test"):
    response = client.post(
        "/api/cases",
        json={
            "project_id": project_id,
            "title": title,
            "steps": [{"action": "Open login", "expected": "Login page appears"}],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def make_cycle(client, project, case):
    plan = client.post(
        "/api/plans",
        json={"project_id": project["id"], "name": "Release test", "release": "1.0", "case_ids": [case["id"]]},
    )
    assert plan.status_code == 200, plan.text
    cycle = client.post(
        "/api/cycles",
        json={"plan_id": plan.json()["id"], "name": "Smoke", "environment": "QA"},
    )
    assert cycle.status_code == 200, cycle.text
    me = client.get("/api/me").json()
    item = client.post(
        f"/api/cycles/{cycle.json()['id']}/cases",
        json={"case_id": case["id"], "assigned_to": me["id"]},
    )
    assert item.status_code == 200, item.text
    return cycle.json(), item.json()


def test_session_cookie_logout_and_versioned_health(client):
    response = client.post("/api/auth/login", json={"email": "nobody@example.test", "password": "bad"})
    assert response.status_code == 401
    assert client.get("/api/me").status_code == 200
    assert client.get("/api/v1/health").json()["api_version"] == "v1"
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/me").status_code == 401


def test_project_access_and_server_side_role_checks(client):
    project = make_project(client)
    admin = client.get("/api/me").json()
    created = client.post(
        "/api/users",
        json={"email": "tester@example.test", "name": "Tester", "password": "tester-password-2026", "role": "tester"},
    )
    assert created.status_code == 200
    assert client.post(
        f"/api/projects/{project['id']}/members",
        json={"user_id": created.json()["id"], "role": "viewer"},
    ).status_code == 200
    case_response = client.post(
        "/api/cases",
        json={"project_id": project["id"], "title": "Admin-only fixture", "steps": []},
    )
    assert case_response.status_code == 200

    outsider = TestClient(main.app)
    login = outsider.post(
        "/api/auth/login",
        json={"email": "tester@example.test", "password": "tester-password-2026"},
    )
    assert login.status_code == 200
    assert outsider.post(
        "/api/cases",
        json={"project_id": project["id"], "title": "Unauthorized"},
    ).status_code == 403
    assert outsider.get(f"/api/cases?project_id={project['id']}").status_code == 200
    assert admin["role"] == "admin"
    outsider.close()


def test_multiple_global_roles_and_test_lead_permissions(client):
    admin_id = client.get("/api/me").json()["id"]
    assert client.patch(
        f"/api/users/{admin_id}",
        json={"roles": ["tester"]},
    ).status_code == 400

    created = client.post(
        "/api/users",
        json={
            "email": "qa-lead@example.test",
            "name": "QA Lead",
            "password": "qa-lead-password-2026",
            "roles": ["admin", "qa_lead"],
        },
    )
    assert created.status_code == 200, created.text
    assert client.patch(
        f"/api/users/{created.json()['id']}",
        json={"roles": ["admin", "qa_lead", "test_lead"], "active": True},
    ).json()["roles"] == ["admin", "qa_lead", "test_lead"]

    created_lead = client.post(
        "/api/users",
        json={
            "email": "test-lead@example.test",
            "name": "Test Lead",
            "password": "test-lead-password-2026",
            "roles": ["test_lead", "tester"],
        },
    )
    assert created_lead.status_code == 200, created_lead.text

    test_lead = TestClient(main.app)
    login = test_lead.post(
        "/api/auth/login",
        json={"email": "test-lead@example.test", "password": "test-lead-password-2026"},
    )
    assert login.status_code == 200
    assert test_lead.get("/api/me").json()["roles"] == ["test_lead", "tester"]
    project = test_lead.post(
        "/api/projects",
        json={"key": "TESTLEAD", "name": "Test Lead Project"},
    )
    assert project.status_code == 200, project.text
    assert test_lead.get("/api/projects").json()[0]["role"] == "lead"
    test_lead.close()


def test_legacy_user_roles_are_migrated(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    with engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE users (id INTEGER PRIMARY KEY, email VARCHAR(255) NOT NULL, "
            "name VARCHAR(120) NOT NULL, password_hash VARCHAR(300) NOT NULL, "
            "role VARCHAR(20), active BOOLEAN)"
        )
        connection.exec_driver_sql(
            "INSERT INTO users (id,email,name,password_hash,role,active) VALUES "
            "(1,'lead@example.test','Lead','hash','lead',1),"
            "(2,'tester@example.test','Tester','hash','tester',1)"
        )

    with engine.begin() as connection:
        main.ensure_database_schema(connection)
    with sessionmaker(bind=engine)() as session:
        assert session.get(main.User, 1).roles == ["qa_lead"]
        assert session.get(main.User, 2).roles == ["tester"]
    engine.dispose()


def test_versioned_cases_independent_attempts_defects_and_reports(client):
    project = make_project(client)
    case = make_case(client, project["id"])
    cycle, item = make_cycle(client, project, case)
    updated = client.put(
        f"/api/cases/{case['id']}",
        json={
            "project_id": project["id"],
            "title": "Updated login test",
            "preconditions": "User is logged out",
            "steps": [{"action": "Use new flow", "expected": "Works"}],
        },
    )
    assert updated.status_code == 200
    snapshot = client.get(f"/api/cycles/{cycle['id']}/cases").json()[0]
    assert snapshot["version"] == 1
    assert snapshot["steps"][0]["action"] == "Open login"

    attempt_ids = []
    for status in ("Failed", "Passed"):
        response = client.post(
            f"/api/cycle-cases/{item['id']}/executions",
            json={
                "status": status,
                "actual_result": f"Observed {status}",
                "notes": "Independent attempt",
                "step_results": [{"step": 0, "status": status, "actual": "Observed"}],
            },
        )
        assert response.status_code == 200, response.text
        attempt_ids.append(response.json()["id"])
    assert len(client.get(f"/api/executions?cycle_id={cycle['id']}").json()) == 2

    defect = client.post(
        "/api/defects",
        json={
            "project_id": project["id"],
            "title": "Login regression",
            "execution_ids": attempt_ids,
        },
    )
    assert defect.status_code == 200, defect.text
    assert sorted(client.get(f"/api/defects/{defect.json()['id']}").json()["execution_ids"]) == sorted(attempt_ids)
    assert client.patch(f"/api/defects/{defect.json()['id']}?status=Retest").status_code == 200
    csv_response = client.get(f"/api/reports/executions.csv?project_id={project['id']}&cycle_id={cycle['id']}")
    assert csv_response.status_code == 200
    assert csv_response.text.count("Independent attempt") == 0
    assert "execution_id,case_id,cycle_id,status" in csv_response.text
    dashboard = client.get(f"/api/dashboard?project_id={project['id']}&release=1.0").json()
    assert dashboard["statuses"]["Passed"] == 1
    assert dashboard["statuses"]["Not Run"] == 0


def test_junit_import_is_reviewable_and_duplicate_safe(client):
    project = make_project(client)
    case = make_case(client, project["id"], "Automated login")
    cycle, _ = make_cycle(client, project, case)
    assert client.post(
        "/api/automation/mappings",
        json={"project_id": project["id"], "external_id": "tests.login.test_valid", "case_id": case["id"]},
    ).status_code == 200
    xml = b"""<testsuites><testsuite><testcase classname="tests.login" name="test_valid"/><testcase classname="tests.login" name="unmapped"/></testsuite></testsuites>"""
    result = client.post(
        f"/api/automation/import/{cycle['id']}",
        files={"file": ("run.xml", xml, "application/xml")},
    )
    assert result.status_code == 200, result.text
    assert result.json()["imported"] == 1
    assert result.json()["unmatched"] == ["tests.login.unmapped"]
    duplicate = client.post(
        f"/api/automation/import/{cycle['id']}",
        files={"file": ("run.xml", xml, "application/xml")},
    )
    assert duplicate.status_code == 409
    assert len(client.get(f"/api/automation/imports?cycle_id={cycle['id']}").json()) == 1


def test_scoped_ci_tokens_evidence_traceability_and_excel(client):
    project = make_project(client)
    case = make_case(client, project["id"])
    cycle, item = make_cycle(client, project, case)

    token_response = client.post(
        f"/api/automation/tokens?project_id={project['id']}&name=CI",
    )
    assert token_response.status_code == 200
    integration_token = token_response.json()["token"]
    xml = b'<testsuite><testcase classname="ci" name="login"/></testsuite>'
    result = client.post(
        f"/api/automation/ci/{cycle['id']}",
        headers={"Authorization": f"Bearer {integration_token}"},
        files={"file": ("pipeline.xml", xml, "application/xml")},
    )
    assert result.status_code == 200, result.text

    execution = client.post(f"/api/cycle-cases/{item['id']}/executions", json={"status": "Passed"})
    evidence = client.post(
        f"/api/executions/{execution.json()['id']}/attachments",
        files={"file": ("evidence.txt", b"test evidence", "text/plain")},
    )
    assert evidence.status_code == 200, evidence.text
    assert client.get(f"/api/attachments/{evidence.json()['id']}").content == b"test evidence"
    rejected = client.post(
        f"/api/executions/{execution.json()['id']}/attachments",
        files={"file": ("unsafe.exe", b"MZ", "application/octet-stream")},
    )
    assert rejected.status_code == 415

    requirement = client.post(
        "/api/requirements",
        json={"project_id": project["id"], "key": "REQ-1", "title": "Login works"},
    )
    assert requirement.status_code == 200, requirement.text
    assert client.post(
        f"/api/requirements/{requirement.json()['id']}/cases?case_id={case['id']}",
    ).status_code == 200
    assert client.get(f"/api/requirements/coverage?project_id={project['id']}").json()[0]["covered"] is True

    field = client.post(
        f"/api/custom-fields?project_id={project['id']}&name=Team&field_type=select&options_json=%5B%22QA%22%5D",
    )
    assert field.status_code == 200, field.text
    bad = client.post(
        "/api/cases",
        json={"project_id": project["id"], "title": "Bad custom value", "custom_values": {"Team": "Other"}},
    )
    assert bad.status_code == 400

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["title", "priority", "kind", "preconditions", "steps"])
    sheet.append(["Imported case", "High", "Manual", "Ready", '[{"action":"Open","expected":"Loaded"}]'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    imported = client.post(
        f"/api/cases/import.xlsx?project_id={project['id']}",
        files={"file": ("cases.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    exported = client.get(f"/api/cases/export/xlsx?project_id={project['id']}")
    assert exported.status_code == 200
    parsed = load_workbook(io.BytesIO(exported.content), read_only=True, data_only=True)
    assert parsed.active.max_row == 3
    assert client.get(f"/api/audit?project_id={project['id']}").status_code == 200


def test_public_registration_persists_tester_and_allows_login(client):
    client.post('/api/auth/logout')
    payload = {'name': ' New Tester ', 'email': 'NEW@example.test', 'password': 'new-tester-password-2026', 'roles': ['admin'], 'role': 'admin'}
    response = client.post('/api/auth/register', json=payload)
    assert response.status_code == 201, response.text
    assert 'password_hash' not in response.json()
    with main.SessionLocal() as session:
        account = session.get(main.User, response.json()['id'])
        assert account.email == 'new@example.test'
        assert account.name == 'New Tester'
        assert account.roles == ['tester']
        assert account.role == 'tester'
        assert account.password_hash != payload['password']
        assert main.verify_password(payload['password'], account.password_hash)
    assert client.post('/api/auth/register', json=payload).status_code == 409
    assert client.post('/api/auth/login', json={'email': 'new@example.test', 'password': payload['password']}).status_code == 200
    assert client.get('/api/me').json()['roles'] == ['tester']
    assert client.get('/api/projects').json() == []
    assert client.post('/api/users', json={'name': 'Other', 'email': 'other@example.test', 'password': 'another-password-2026'}).status_code == 403


@pytest.mark.parametrize('changes,expected', [
    ({'name': ''}, 422),
    ({'name': '   '}, 400),
    ({'email': 'not-an-email'}, 400),
    ({'password': 'short'}, 422),
])
def test_public_registration_rejects_invalid_input(client, changes, expected):
    client.post('/api/auth/logout')
    payload = {'name': 'Tester', 'email': 'invalid@example.test', 'password': 'valid-password-2026'}
    payload.update(changes)
    assert client.post('/api/auth/register', json=payload).status_code == expected
    with main.SessionLocal() as session:
        assert session.scalar(main.select(main.func.count()).select_from(main.User)) == 1


def test_dashboard_router_aliases_filters_and_permissions(client):
    project = make_project(client)
    case = make_case(client, project['id'])
    for title, actual_result in (
        ('Spreadsheet pass', 'PASS'),
        ('Spreadsheet pass lowercase', ' pass '),
        ('Spreadsheet fail', 'FAIL'),
        ('Spreadsheet in progress', 'Pending'),
        ('Spreadsheet skip', 'Skipped'),
        ('Spreadsheet block', 'BLOCK'),
        ('No recorded result', ''),
    ):
        response = client.post('/api/cases', json={
            'project_id': project['id'],
            'title': title,
            'import_metadata': {'actual_result': actual_result},
        })
        assert response.status_code == 200, response.text
    cycle, item = make_cycle(client, project, case)
    path = f"/api/dashboard?project_id={project['id']}"
    initial = client.get(path)
    assert initial.status_code == 200
    assert initial.json()['statuses']['Not Run'] == 1
    assert initial.json()['case_statuses'] == {
        'Pass': 2,
        'Fail': 1,
        'In Progress': 1,
        'Skip': 1,
        'Block': 1,
    }
    assert initial.json()['progress'] == 0
    assert client.get(path.replace('/api/', '/api/v1/')).json() == initial.json()
    for status in ('Failed', 'Passed'):
        assert client.post(f"/api/cycle-cases/{item['id']}/executions", json={'status': status}).status_code == 200
    result = client.get(path + f"&cycle_id={cycle['id']}&release=1.0").json()
    assert result['executions'] == 2
    assert result['statuses']['Passed'] == 1
    assert result['progress'] == 1
    assert client.get(path + '&release=missing').json()['cycles'] == 0
    other = client.post('/api/projects', json={'key': 'OTHER', 'name': 'Other'}).json()
    assert client.get(f"/api/dashboard?project_id={other['id']}&cycle_id={cycle['id']}").status_code == 400
    client.post('/api/auth/logout')
    assert client.get(path).status_code == 401
    account = {'name': 'Dashboard outsider', 'email': 'dashboard-outsider@example.test', 'password': 'dashboard-password-2026'}
    assert client.post('/api/auth/register', json=account).status_code == 201
    assert client.post('/api/auth/login', json=account).status_code == 200
    assert client.get(path).status_code == 403



def test_case_excel_import_summary_and_display_data(client):
    project = make_project(client)
    book = Workbook()
    sheet = book.active
    sheet.append(['title', 'priority', 'kind', 'preconditions', 'steps'])
    sheet.append(['Imported login', 'High', 'Manual', 'Logged out', '[{"action":"Open login","expected":"Page appears"}]'])
    sheet.append(['', 'Low', 'Manual', '', '[]'])
    sheet.append(['Invalid steps', 'Low', 'Manual', '', 'not json'])
    sheet.append(['Invalid priority', 'Wrong', 'Manual', '', '[]'])
    sheet.append(['Invalid step object', 'Low', 'Manual', '', '[1]'])
    buffer = io.BytesIO()
    book.save(buffer)
    result = client.post(f"/api/v1/cases/import.xlsx?project_id={project['id']}", files={'file': ('cases.xlsx', buffer.getvalue())})
    assert result.status_code == 200, result.text
    summary = result.json()
    assert summary['imported'] == 1
    assert [row['row'] for row in summary['errors']] == [3, 4, 5, 6]
    rows = client.get(f"/api/cases?project_id={project['id']}").json()
    assert len(rows) == 1
    assert rows[0]['id'] == summary['case_ids'][0]
    assert rows[0]['title'] == 'Imported login'
    details = client.get(f"/api/cases/{rows[0]['id']}").json()
    assert details['preconditions'] == 'Logged out'
    assert details['steps'][0]['expected'] == 'Page appears'
    assert rows[0]['preconditions'] == details['preconditions']
    assert rows[0]['steps'] == details['steps']
    # A viewer can read the imported cases but cannot import more.
    user = client.post('/api/users', json={'name':'Reader','email':'excel-reader@example.test','password':'excel-reader-password-2026'}).json()
    client.post(f"/api/projects/{project['id']}/members", json={'user_id':user['id'],'role':'viewer'})
    client.post('/api/auth/logout')
    client.post('/api/auth/login', json={'email':'excel-reader@example.test','password':'excel-reader-password-2026'})
    assert client.get(f"/api/cases?project_id={project['id']}").status_code == 200
    assert client.post(f"/api/cases/import.xlsx?project_id={project['id']}", files={'file': ('cases.xlsx', buffer.getvalue())}).status_code == 403


def test_case_list_is_ordered_oldest_first(client):
    project = make_project(client)
    created_ids = [make_case(client, project["id"])["id"] for _ in range(3)]

    rows = client.get(f"/api/cases?project_id={project['id']}").json()

    assert [row["id"] for row in rows] == created_ids


def test_case_create_and_edit_store_spreadsheet_columns(client):
    project = make_project(client)
    payload = {
        "project_id": project["id"],
        "title": "Verify pipeline loads",
        "priority": "High",
        "kind": "Manual",
        "status": "Active",
        "preconditions": "Admin is logged in",
        "steps": [{"action": "Open pipeline", "expected": "Pipeline is displayed"}],
        "import_metadata": {
            "feature": "Radiologist Pipeline",
            "test_type": "UI/Functional",
            "actual_result": "Pipeline displayed",
            "sheet_status": "Passed",
            "expected_result": "Pipeline is displayed",
            "test_steps": "Open pipeline"
        }
    }

    created = client.post("/api/cases", json=payload)
    assert created.status_code == 200, created.text
    case_id = created.json()["id"]
    details = client.get(f"/api/cases/{case_id}").json()
    assert details["import_metadata"] == payload["import_metadata"]

    payload["import_metadata"]["actual_result"] = "Pipeline displayed without errors"
    updated = client.put(f"/api/cases/{case_id}", json=payload)
    assert updated.status_code == 200, updated.text
    assert client.get(f"/api/cases/{case_id}").json()["import_metadata"] == payload["import_metadata"]



def test_spreadsheet_columns_plain_steps_and_metadata_roundtrip(client):
    project = make_project(client)
    book = Workbook()
    sheet = book.active
    sheet.append(['-', 'Feature', 'Test Scenario', 'Preconditions', 'Test Steps', 'Expected Result', 'Priority', 'Test Type', 'Actual Result', 'Status'])
    sheet.append(['RAD-TC-001', 'Radiologist Pipeline', 'Verify pipeline loads', 'Admin logged in', '1. Open pipeline.\n2. Observe cards.', 'Cards appear in correct stages.', 'critical', 'UI/Functional', 'PASS', 'Passed'])
    sheet.append(['RAD-TC-002', 'Pipeline Count', 'Verify total count', 'Records exist', 'Observe total count.', 'Total matches cards.', 'High', 'Functional', '', ''])
    buffer = io.BytesIO()
    book.save(buffer)
    imported = client.post(f"/api/cases/import.xlsx?project_id={project['id']}", files={'file': ('pipeline.xlsx', buffer.getvalue())})
    assert imported.status_code == 200, imported.text
    assert imported.json()['imported'] == 2
    assert imported.json()['errors'] == []
    case_id = imported.json()['case_ids'][0]
    details = client.get(f'/api/cases/{case_id}').json()
    assert details['title'] == 'Verify pipeline loads'
    assert details['priority'] == 'Critical'
    assert details['kind'] == 'Manual'
    assert details['status'] == 'Active'
    assert details['steps'] == [{'action':'Open pipeline.','expected':''}, {'action':'Observe cards.','expected':'Cards appear in correct stages.'}]
    metadata = details['import_metadata']
    assert metadata['source_id'] == 'RAD-TC-001'
    assert metadata['test_type'] == 'UI/Functional'
    assert metadata['actual_result'] == 'PASS'
    assert metadata['sheet_status'] == 'Passed'
    assert client.get(f"/api/dashboard?project_id={project['id']}").json()['executions'] == 0
    updated = client.put(f'/api/cases/{case_id}', json={**details, 'title':'Updated title'})
    assert updated.status_code == 200, updated.text
    assert updated.json()['import_metadata'] == metadata
    exported = client.get(f"/api/cases/export/xlsx?project_id={project['id']}")
    other = client.post('/api/projects', json={'key':'ROUNDTRIP','name':'Roundtrip'}).json()
    reimported = client.post(f"/api/cases/import.xlsx?project_id={other['id']}", files={'file': ('roundtrip.xlsx', exported.content)})
    assert reimported.status_code == 200, reimported.text
    assert reimported.json()['imported'] == 2
    copy = client.get(f"/api/cases/{reimported.json()['case_ids'][0]}").json()
    assert copy['import_metadata'] == metadata
    assert copy['steps'] == details['steps']


def test_import_metadata_migration_preserves_existing_rows(tmp_path):
    from alembic import command
    from alembic.config import Config
    database_url = f"sqlite:///{tmp_path / 'before_metadata.db'}"
    engine = create_engine(database_url)
    with engine.begin() as connection:
        connection.exec_driver_sql('CREATE TABLE test_cases (id INTEGER PRIMARY KEY, title TEXT)')
        connection.exec_driver_sql("INSERT INTO test_cases VALUES (1, 'Existing case')")
        connection.exec_driver_sql('CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL)')
        connection.exec_driver_sql("INSERT INTO alembic_version VALUES ('0001_srs_schema')")
    config = Config('alembic.ini')
    config.set_main_option('sqlalchemy.url', database_url)
    command.upgrade(config, 'head')
    with engine.connect() as connection:
        row = connection.exec_driver_sql('SELECT title, import_metadata FROM test_cases WHERE id=1').one()
        assert row == ('Existing case', '{}')
    engine.dispose()
